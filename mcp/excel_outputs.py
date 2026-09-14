"""Manual QA Excel outputs — iPay matrix + Book1 visual (gold LATEST architecture)."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from paths import REPO_ROOT, WORKSPACE_ROOT
except ImportError:
    REPO_ROOT = Path(__file__).resolve().parent.parent
    WORKSPACE_ROOT = REPO_ROOT / "workspace"

try:
    from book1_contract import (
        BOOK1_COLUMNS,
        BOOK1_HEADER_FILL,
        BOOK1_SHEET_NAME,
        BOOK1_WIDTHS,
        contract_summary,
        ensure_full_book1_row,
        ensure_full_what_is_testing,
        simple_explanation,
        validate_book1_workbook,
    )
except ImportError:
    from mcp.book1_contract import (  # type: ignore
        BOOK1_COLUMNS,
        BOOK1_HEADER_FILL,
        BOOK1_SHEET_NAME,
        BOOK1_WIDTHS,
        contract_summary,
        ensure_full_book1_row,
        ensure_full_what_is_testing,
        simple_explanation,
        validate_book1_workbook,
    )

IPAY_COLUMNS = [
    "Area",
    "Concern",
    "User story",
    "Status",
    "Change made?",
    "Change / verification notes (English)",
    "Commit / cycle",
]

SAMPLE = [
    ("Access", "Entry API 401 blocks grid.", "Fail", "No", "Proof path here"),
    ("List", "Search returns expected rows.", "Pass", "No", "Proof path here"),
    ("Create", "Create control missing.", "Fail", "No", "AC blocked"),
    ("Validation", "Empty search validated.", "Pass", "No", "Toast shown"),
    ("Detail", "AC column missing on view.", "Fail", "No", "Log bug"),
    ("RBAC", "Checker blocked.", "Blocked", "No", "Need grant"),
]

AREAS = [
    "Access",
    "Navigation",
    "List",
    "Search",
    "Create",
    "Edit",
    "View",
    "Submit",
    "Validation",
    "Empty",
    "Error",
    "Pagination",
    "RBAC",
    "API",
    "Regression",
]


def output_dirs():
    ws = WORKSPACE_ROOT
    reports, artifacts = ws / "reports", REPO_ROOT / "artifacts"
    outputs = ws / "outputs"
    dl = Path.home() / "Downloads"
    for p in (reports, artifacts, dl, outputs, reports / "book1-per-story"):
        p.mkdir(parents=True, exist_ok=True)
    return {
        "workspace_reports": str(reports.resolve()),
        "workspace_outputs": str(outputs.resolve()),
        "repo_artifacts": str(artifacts.resolve()),
        "downloads": str(dl.resolve()),
    }


def access_guide():
    d = REPO_ROOT / "docs"
    return {
        "ok": True,
        "repo_root": str(REPO_ROOT.resolve()),
        "workspace_root": str(WORKSPACE_ROOT.resolve()),
        "flow_doc": str((d / "MANUAL-QA-FLOW-FULL.md").resolve()),
        "flow_chart": str((d / "FLOW-CHART.md").resolve()),
        "excel_doc": str((d / "EXCEL-OUTPUTS.md").resolve()),
        "outputs": output_dirs(),
        "ipay_columns": IPAY_COLUMNS,
        "book1_columns": BOOK1_COLUMNS,
        "book1_contract": contract_summary(),
        "target_rows_per_story": 110,
        "excel_tools": [
            "manualqa_status",
            "manualqa_generate_sample_excel",
            "manualqa_generate_book1_sample",
            "manualqa_book1_append_row",
            "manualqa_validate_book1",
            "manualqa_generate_from_result",
            "manualqa_list_outputs",
        ],
        "manual_tools": [
            "humanize_screenshot",
            "humanize_capture",
            "humanize_click",
            "humanize_type",
            "humanize_browser_open",
        ],
        "flow_tools": [
            "manualqa_boot",
            "manualqa_laws",
            "manualqa_todo_list",
            "manualqa_complete_step",
            "manualqa_monitor_heartbeat",
        ],
    }


def _rows(story_key, cycle, min_rows):
    rows = [[a, c, story_key, s, ch, n, cycle] for a, c, s, ch, n in SAMPLE]
    i = 0
    while len(rows) < min_rows:
        a = AREAS[i % len(AREAS)]
        rows.append(
            [
                a,
                f"[Template {len(rows)+1}] Define failure risk for {a}.",
                story_key,
                "Pending",
                "No",
                "Replace after execution.",
                cycle,
            ]
        )
        i += 1
    return rows


def _style(ws):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    hf = PatternFill("solid", fgColor="1F4E79")
    wf = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(*(Side(style="thin"),) * 4)
    for col, name in enumerate(IPAY_COLUMNS, 1):
        c = ws.cell(1, col, name)
        c.fill, c.font, c.alignment, c.border = hf, wf, wrap, thin
    for i, w in enumerate([22, 72, 14, 28, 14, 88, 42], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def generate_ipay_sample(story_key="STORY-001", story_title="Example", min_rows=110, cycle=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return {"ok": False, "error": "pip install openpyxl"}
    cycle = cycle or datetime.now(timezone.utc).strftime("%Y-%m-%d sample")
    rows = _rows(story_key, cycle, max(min_rows, 110))
    wb = Workbook()
    cov = wb.active
    cov.title = "Cover"
    cov["A1"] = f"{story_key} manual QA matrix"
    cov["A1"].font = Font(bold=True, size=14)
    cov["A3"], cov["B3"] = "Rows", len(rows)
    ws = wb.create_sheet("Matrix")
    ws.append(IPAY_COLUMNS)
    _style(ws)
    for r in rows:
        ws.append(r)
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    dirs = output_dirs()
    fn = f"{story_key}-manual-qa-matrix.xlsx"
    paths = [
        Path(dirs["workspace_reports"]) / fn,
        Path(dirs["workspace_outputs"]) / story_key / fn,
        Path(dirs["repo_artifacts"]) / fn,
        Path(dirs["downloads"]) / fn,
    ]
    written = []
    for p in paths:
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(p)
        written.append(str(p.resolve()))
    return {
        "ok": True,
        "story_key": story_key,
        "rows": len(rows),
        "columns": IPAY_COLUMNS,
        "files": written,
        "primary": written[0],
    }


def _book1_path(story_key: str) -> Path:
    dirs = output_dirs()
    p = Path(dirs["workspace_outputs"]) / story_key / f"{story_key}-Book1.xlsx"
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def generate_book1_sample(story_key="STORY-001"):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return {"ok": False, "error": "pip install openpyxl"}

    wb = Workbook()
    ws = wb.active
    ws.title = BOOK1_SHEET_NAME
    header = PatternFill("solid", fgColor=BOOK1_HEADER_FILL)
    hf = Font(bold=True, color="FFFFFF")
    for col, name in enumerate(BOOK1_COLUMNS, 1):
        cell = ws.cell(1, col, name)
        cell.fill = header
        cell.font = hf
    sample_rows = [
        [
            f"{story_key} | FusionX / GUI",
            "Menu exists or honest N/A proof. — PASSED",
            f"reports/proof/{story_key}/01.png",
            f"TC-{story_key}-01: Menu exists or honest N/A proof.",
            "Predicted pass; record actual vs prediction in English.",
            "PASS — NO PRODUCT BUG\n\n2nd QA matched.",
            simple_explanation(
                issue="Menu exists or honest N/A proof. — PASSED",
                what_is_testing=f"TC-{story_key}-01: Menu exists or honest N/A proof.",
                result="PASS",
            ),
        ],
        [
            f"{story_key} | FusionX / GUI",
            "Required field missing on save. — FAILED",
            f"reports/proof/{story_key}/02.png",
            f"TC-{story_key}-02: Required field missing on save.",
            "Predicted ready; actual=Fail on Kenya UAT.",
            "REAL PRODUCT ISSUE — 2nd QA CONFIRMED",
            simple_explanation(
                issue="Required field missing on save. — FAILED",
                what_is_testing=f"TC-{story_key}-02: Required field missing on save.",
                result="FAIL",
            ),
        ],
    ]
    for r in sample_rows:
        ws.append(r)
        for c in range(1, 8):
            ws.cell(ws.max_row, c).alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(BOOK1_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22

    dirs = output_dirs()
    primary = _book1_path(story_key)
    copies = [
        primary,
        Path(dirs["workspace_reports"]) / "book1-per-story" / primary.name,
        Path(dirs["downloads"]) / primary.name,
        Path(dirs["repo_artifacts"]) / primary.name,
    ]
    written = []
    for p in copies:
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(p)
        written.append(str(p.resolve()))
    return {
        "ok": True,
        "columns": BOOK1_COLUMNS,
        "files": written,
        "primary": written[0],
        "note": "Embed real PNGs in Screenshot column; red-fill Issue cell for bugs; crop proof to bug point.",
    }


def book1_append_row(
    story_key: str,
    area: str,
    issue: str,
    screenshot_path: str,
    what_is_testing: str,
    why_prediction: str,
    second_qa: str,
    english_explanation: str = "",
    is_bug: bool = False,
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, PatternFill
    except ImportError:
        return {"ok": False, "error": "pip install openpyxl"}

    path = _book1_path(story_key)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
        # migrate old 6-col books to 7-col if needed
        if ws.cell(1, 7).value is None:
            ws.cell(1, 7, BOOK1_COLUMNS[6])
    else:
        gen = generate_book1_sample(story_key)
        if not gen.get("ok"):
            return gen
        wb = load_workbook(path)
        ws = wb.active
        # clear sample data rows keep header
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

    full = ensure_full_book1_row(
        area=area,
        issue=issue,
        what_is_testing=what_is_testing,
        why_prediction=why_prediction,
        second_qa=second_qa,
        english_explanation=english_explanation,
        story_key=story_key,
        is_bug=is_bug,
    )
    ws.append(
        [
            full["area"],
            full["issue"],
            screenshot_path,
            full["what_is_testing"],
            full["why_prediction"],
            full["second_qa"],
            full["english_explanation"],
        ]
    )
    row = ws.max_row
    for c in range(1, 8):
        ws.cell(row, c).alignment = Alignment(wrap_text=True, vertical="top")
    if is_bug:
        ws.cell(row, 2).fill = PatternFill("solid", fgColor="FF6B6B")
    # try embed image (BytesIO — openpyxl closes file handles on multi-save)
    shot = Path(screenshot_path)
    if shot.exists() and shot.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        try:
            from io import BytesIO

            from PIL import Image as PILImage

            buf = BytesIO()
            with PILImage.open(shot) as pil:
                pil = pil.convert("RGB")
                pil.thumbnail((320, 180))
                pil.save(buf, format="PNG")
            buf.seek(0)
            img = XLImage(buf)
            img.width = 320
            img.height = 180
            ws.add_image(img, f"C{row}")
            ws.row_dimensions[row].height = max(120, ws.row_dimensions[row].height or 15)
        except Exception as e:
            ws.cell(row, 3).value = f"{screenshot_path} (embed failed: {e})"

    dirs = output_dirs()
    primary = path
    primary.parent.mkdir(parents=True, exist_ok=True)
    wb.save(primary)
    written = [str(primary.resolve())]
    import shutil

    for p in (
        Path(dirs["workspace_reports"]) / "book1-per-story" / path.name,
        Path(dirs["downloads"]) / path.name,
    ):
        try:
            p.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(primary, p)
            written.append(str(p.resolve()))
        except Exception:
            pass
    return {"ok": True, "row": row, "is_bug": is_bug, "files": written, "columns": BOOK1_COLUMNS}


def list_outputs(limit=50):
    found = []
    for label, folder in output_dirs().items():
        root = Path(folder)
        if not root.exists():
            continue
        for p in sorted(root.rglob("*.xlsx"), key=lambda x: x.stat().st_mtime, reverse=True):
            found.append({"location": label, "path": str(p.resolve())})
            if len(found) >= limit:
                break
        if len(found) >= limit:
            break
    return {"ok": True, "count": len(found), "files": found, "output_dirs": output_dirs()}


def validate_book1(path: str) -> dict[str, Any]:
    """MCP-facing: validate a Book1 xlsx against gold LATEST + screenshot coverage; sync Gatekeeper."""
    result = validate_book1_workbook(path)
    try:
        from gatekeeper_bridge import sync_book1_validate

        result["gatekeeper"] = sync_book1_validate(result)
    except Exception as e:  # noqa: BLE001
        result["gatekeeper"] = {"ok": False, "skipped": True, "error": str(e)}
    return result


def generate_from_result(story_key):
    rp = WORKSPACE_ROOT / "reports" / "live-manual" / f"{story_key}-RESULT.md"
    if not rp.exists():
        return {
            "ok": False,
            "error": str(rp),
            "hint": "Execute headed QA first, then generate.",
        }
    return generate_ipay_sample(story_key, cycle="from RESULT")


def testcase_sufficiency(task_key: str, new_count: int, existing_count: int, notes: str) -> dict[str, Any]:
    """Step 7c — agent must iterate until satisfied."""
    from workspace_contract import load_state, save_state

    enough = new_count >= 20 and (new_count + existing_count) >= 40
    result = {
        "ok": True,
        "task_key": task_key,
        "new_test_cases": new_count,
        "existing_test_cases": existing_count,
        "total": new_count + existing_count,
        "enough": enough,
        "recommendation": (
            "Satisfied baseline (≥20 new, ≥40 total). Still deepen risk areas if map shows gaps."
            if enough
            else "NOT enough — keep generating until coverage matches map + stories. Re-run cycle."
        ),
        "notes": notes,
    }
    st = load_state()
    suf = st.setdefault("sufficiency", {"cycles": []})
    suf["enough"] = enough
    suf["last"] = result
    suf.setdefault("cycles", []).append(result)
    suf["cycles"] = suf["cycles"][-50:]
    save_state(st)
    return result
